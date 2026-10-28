"""
Reading the two Cisco Learning Credits exports.

Both are single-sheet workbooks with a header row, so the work here is not
the XLSX but the shape of what Cisco puts in it:

  * One registration can be **split across several SO numbers** when a single
    pool cannot cover it. Those arrive as separate rows sharing one
    `TransactionId` — transaction 1081935 is one 43-credit IT Leadership pass
    debited 24 from one SO and 19 from another. Counting rows would count
    that person twice.

  * `DebitStatus` carries the outcome. `SEND_MANUAL` and `DEBITED(Approved)`
    are live registrations; `DEBITED(Denied)` and
    `DEBITED(Canceled - LP Refunded)` are not, and a `REFUND` row gives the
    credits back.

  * `CourseAcronym` carries the edition and the pass type in one string, and
    the spelling changed between years: `CLEMEA2027-ITL-EB`,
    `CLEMEA2026-FC-STD`, `CLEMEA2025FCST`, `CLSD2025EXPGPST`. A registration
    with add-ons holds a comma-joined list.

  * `Attendees` is the person ("Alex Moreau (alex.moreau@example.com)");
    `InitiatedBy` is whoever booked, which is not always the same.
"""

from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# 'CLEMEA2027-ITL-EB' -> ('CLEMEA', '2027', '-ITL-EB')
#       'CLEMEA2025FCST' -> ('CLEMEA', '2025', 'FCST')
_ACRONYM = re.compile(r"^(CL[A-Z]*?)(\d{4})(.*)$", re.IGNORECASE)

# 'Alex Moreau (alex.moreau@example.com)' -> name, email
_ATTENDEE = re.compile(r"^\s*(.*?)\s*\(([^)]*@[^)]*)\)\s*$")

LIVE_DEBIT_STATUSES = {"SEND_MANUAL", "DEBITED(APPROVED)", "REISSUED"}
VOID_DEBIT_STATUSES = {"DEBITED(DENIED)", "DEBITED(CANCELED - LP REFUNDED)"}


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _as_number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def read_sheet(path: Path) -> list[dict[str, Any]]:
    """Every row of the first sheet as a dict keyed by the header row.

    `read_only` keeps a large export off the heap, and `data_only` gives the
    cached result of a formula rather than its text.
    """
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        keys = [_as_text(h) for h in header]
        return [
            dict(zip(keys, row))
            for row in rows
            # A trailing all-empty row is common in these exports.
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()


def parse_acronym(acronym: str) -> tuple[str | None, str]:
    """Split a course acronym into its edition code and the rest.

    Returns ('CLEMEA2027', 'ITL-EB') for 'CLEMEA2027-ITL-EB'. The edition is
    None for anything that is not a Cisco Live product, which is most of the
    file — exam vouchers and instructor-led courses share the same report.
    """
    match = _ACRONYM.match(acronym.strip())
    if not match:
        return None, ""
    family, year, tail = match.groups()
    return f"{family.upper()}{year}", tail.strip(" -").upper()


def parse_attendee(raw: str) -> tuple[str, str]:
    """Pull the display name and the email out of the `Attendees` cell."""
    match = _ATTENDEE.match(raw)
    if match:
        return match.group(1).strip(), match.group(2).strip().lower()
    # Some rows carry a bare email and no name.
    if "@" in raw:
        return raw.strip(), raw.strip().lower()
    return raw.strip(), ""


# --------------------------------------------------------------------------
# Transaction report
# --------------------------------------------------------------------------


def read_customer(path: Path) -> str:
    """The account the export belongs to, from its `CustomerName` column.

    Whose console this is, is a property of the data, not of the source. The
    most common value wins rather than the first, so one stray blank or
    mistyped row cannot rename the account in the header.
    """
    names = Counter(
        _as_text(row.get("CustomerName")) for row in read_sheet(path) if row.get("CustomerName")
    )
    return names.most_common(1)[0][0] if names else ""


def read_transactions(path: Path) -> list[dict[str, Any]]:
    """Cisco Live rows of the transaction report, one dict per row.

    Non-Cisco-Live rows (exam vouchers, ILT courses, the credit purchases
    themselves) are dropped here: `PartnerName` is what distinguishes them,
    and every caller downstream only ever wants the event rows.
    """
    records: list[dict[str, Any]] = []

    for row in read_sheet(path):
        partner = _as_text(row.get("PartnerName")).upper()
        if "CISCO LIVE" not in partner:
            continue

        acronyms = [a.strip() for a in _as_text(row.get("CourseAcronym")).split(",") if a.strip()]
        editions = [parse_acronym(a) for a in acronyms]
        # A registration with add-ons lists several products from the same
        # edition, so the first non-empty edition speaks for the row.
        edition = next((code for code, _ in editions if code), None)
        suffixes = [tail for _, tail in editions if tail]

        name, email = parse_attendee(_as_text(row.get("Attendees")))
        status = _as_text(row.get("DebitStatus")).upper()
        kind = _as_text(row.get("TransactionType")).upper()

        records.append(
            {
                "transactionId": _as_text(row.get("TransactionId")),
                "date": _as_date(row.get("TransactionDate")),
                "type": kind,
                "status": status,
                "so": _as_text(row.get("SONumber")),
                "credits": _as_number(row.get("Credits")),
                "edition": edition,
                "suffixes": suffixes,
                "acronyms": acronyms,
                "courseName": _as_text(row.get("CourseName")),
                "country": _as_text(row.get("ServiceCountry")),
                "attendeeName": name,
                "attendeeEmail": email,
                "initiatedBy": _as_text(row.get("InitiatedBy")),
                # Live if it is an approved debit; void if denied or
                # cancelled; a refund is neither, it reverses an earlier one.
                "live": kind == "DEBIT" and status in LIVE_DEBIT_STATUSES,
                "void": kind == "DEBIT" and status in VOID_DEBIT_STATUSES,
                "refund": kind == "REFUND",
            }
        )

    return records


# --------------------------------------------------------------------------
# SO history
# --------------------------------------------------------------------------


def read_credit_pools(path: Path, today: date | None = None) -> list[dict[str, Any]]:
    """One dict per SO: what was bought, what is left, and when it dies.

    `DaysRemaining` is in the export but it was computed on the day of the
    export, so it is recomputed here against today instead of trusted.
    """
    today = today or date.today()
    pools: list[dict[str, Any]] = []

    for row in read_sheet(path):
        so = _as_text(row.get("SONumber"))
        if not so:
            continue

        purchased = _as_number(row.get("Purchased"))
        balance = _as_number(row.get("Balance"))
        expired = _as_number(row.get("Expired"))
        expiry = _as_date(row.get("ExpirationDate"))

        pools.append(
            {
                "so": so,
                "purchased": purchased,
                "balance": balance,
                "expired": expired,
                # Whatever is neither still available nor written off has
                # been spent. The export has no "used" column of its own.
                "used": purchased - balance - expired,
                "purchaseDate": _as_date(row.get("PurchaseDate")),
                "expirationDate": expiry,
                "daysRemaining": (expiry - today).days if expiry else None,
            }
        )

    # Pools that still hold credits first, soonest deadline at the top —
    # those are the ones there is anything to do about. Exhausted and
    # expired pools follow, newest first, as history rather than as the
    # first thing read.
    pools.sort(
        key=lambda p: (
            p["balance"] <= 0,
            (p["expirationDate"] or date.max) if p["balance"] > 0 else date.max,
            -(p["expirationDate"] or date.min).toordinal() if p["expirationDate"] else 0,
        )
    )
    return pools
